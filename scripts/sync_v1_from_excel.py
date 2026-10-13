"""从最新 Excel 的新疆黄色单元格生成 V1 指标清单。

只读取原始 Excel。人工核验状态来自 data/v1_review_metadata.csv，
同步不会因 Excel 出现数值而自动提升核验状态。
"""

from __future__ import annotations

import argparse
import colorsys
import csv
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SEARCH_DIR = PROJECT_ROOT / "imports"
DEFAULT_REVIEW = PROJECT_ROOT / "data" / "v1_review_metadata.csv"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "v1_xinjiang_metrics.csv"
SHEET_NAME = "省级数据录入"
DICTIONARY_SHEET = "指标字典"
EXPECTED_COUNT = 6


def find_latest_excel(search_dir: Path) -> Path:
    candidates = [
        path
        for path in search_dir.glob("*.xlsx")
        if not path.name.startswith("~$")
    ]
    if not candidates:
        raise FileNotFoundError(f"没有在 {search_dir} 找到 Excel 文件")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def find_province_row(sheet, aliases: set[str]) -> int:
    matches = [
        row
        for row in range(1, sheet.max_row + 1)
        if str(sheet.cell(row, 1).value or "").strip() in aliases
    ]
    if len(matches) != 1:
        raise ValueError(f"新疆行匹配数量异常: {matches}")
    return matches[0]


def theme_colors(workbook) -> list[str]:
    if not workbook.loaded_theme:
        return []
    root = ElementTree.fromstring(workbook.loaded_theme)
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", namespace)
    colors: list[str] = []
    if scheme is None:
        return colors
    for item in list(scheme):
        child = next(iter(item), None)
        if child is None:
            colors.append("")
        else:
            colors.append(child.attrib.get("val") or child.attrib.get("lastClr") or "")
    return colors


def apply_tint(rgb: str, tint: float) -> str:
    channels = [int(rgb[index : index + 2], 16) for index in (0, 2, 4)]
    adjusted = []
    for value in channels:
        result = value * (1 + tint) if tint < 0 else value + (255 - value) * tint
        adjusted.append(max(0, min(255, round(result))))
    return "".join(f"{value:02X}" for value in adjusted)


def resolve_fill_rgb(cell, workbook) -> tuple[str | None, str, str]:
    fill = cell.fill
    color = fill.fgColor
    raw = ""
    resolved: str | None = None
    if color.type == "rgb" and isinstance(color.rgb, str):
        raw = color.rgb.upper()
        resolved = raw[-6:]
    elif color.type == "indexed" and isinstance(color.indexed, int):
        raw = str(color.indexed)
        if 0 <= color.indexed < len(COLOR_INDEX):
            resolved = COLOR_INDEX[color.indexed].upper()[-6:]
    elif color.type == "theme" and isinstance(color.theme, int):
        raw = f"theme={color.theme};tint={color.tint}"
        colors = theme_colors(workbook)
        if 0 <= color.theme < len(colors) and colors[color.theme]:
            resolved = apply_tint(colors[color.theme].upper()[-6:], float(color.tint or 0))
    return resolved, str(color.type), raw


def is_yellow(cell, workbook) -> tuple[bool, str | None, str, str]:
    rgb, color_type, raw = resolve_fill_rgb(cell, workbook)
    if cell.fill.fill_type not in {"solid", "patternSolid"} or not rgb:
        return False, rgb, color_type, raw
    red, green, blue = (int(rgb[index : index + 2], 16) / 255 for index in (0, 2, 4))
    hue, saturation, value = colorsys.rgb_to_hsv(red, green, blue)
    hue_degrees = hue * 360
    yellow = 42 <= hue_degrees <= 68 and saturation >= 0.25 and value >= 0.65
    return yellow, rgb, color_type, raw


def parse_unit(unit_note: Any) -> str:
    return str(unit_note or "").split("；", 1)[0].strip()


def load_review_metadata(path: Path) -> dict[str, dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    return {row["excel_header"].strip(): row for row in rows}


def dictionary_by_name(workbook) -> dict[str, dict[str, Any]]:
    sheet = workbook[DICTIONARY_SHEET]
    headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    return {
        str(sheet.cell(row, 2).value).strip(): dict(
            zip(headers, [sheet.cell(row, col).value for col in range(1, sheet.max_column + 1)])
        )
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, 2).value
    }


def extract_metrics(excel_path: Path, review_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(excel_path, data_only=False, read_only=False)
    values_workbook = load_workbook(excel_path, data_only=True, read_only=False)
    sheet = workbook[SHEET_NAME]
    values_sheet = values_workbook[SHEET_NAME]
    row = find_province_row(sheet, {"新疆", "新疆维吾尔自治区"})
    review = load_review_metadata(review_path)
    dictionary = dictionary_by_name(workbook)
    metrics: list[dict[str, Any]] = []

    for col in range(4, sheet.max_column):
        cell = sheet.cell(row, col)
        yellow, rgb, color_type, raw_color = is_yellow(cell, workbook)
        if not yellow:
            continue
        name = str(sheet.cell(1, col).value or "").strip()
        metadata = review.get(name)
        if metadata is None:
            raise ValueError(f"黄色指标“{name}”缺少人工审核元数据，停止同步")
        dictionary_row = dictionary.get(name)
        if dictionary_row is None:
            raise ValueError(f"黄色指标“{name}”未在指标字典登记，停止同步")
        unit = parse_unit(sheet.cell(4, col).value)
        if unit != str(dictionary_row.get("单位") or "").strip():
            raise ValueError(
                f"{cell.coordinate} {name} 单位变化: 表头={unit}, 指标字典={dictionary_row.get('单位')}"
            )
        year = sheet.cell(row + 1, col).value or sheet.cell(row, 2).value
        value = values_sheet.cell(row, col).value
        if value is None:
            value = cell.value
        metrics.append(
            {
                "province_name": "新疆维吾尔自治区",
                "province_code": str(sheet.cell(row, 3).value).zfill(6),
                "source_cell": cell.coordinate,
                "source_column": get_column_letter(col),
                "metric_key": metadata["metric_key"],
                "metric_name": name,
                "value": value,
                "unit": unit,
                "year": int(year),
                "verification_status": metadata["verification_status"],
                "source_name": metadata["source_name"],
                "source_url": metadata["source_url"],
                "verification_note": metadata["verification_note"],
                "fill_type": cell.fill.fill_type,
                "color_type": color_type,
                "color_value": rgb or raw_color,
                "excel_file": excel_path.name,
                "display_order": len(metrics) + 1,
            }
        )

    if len(metrics) != EXPECTED_COUNT:
        cells = [metric["source_cell"] for metric in metrics]
        raise ValueError(f"黄色指标数量为 {len(metrics)}，预期6；已识别单元格: {cells}")
    if len({metric["metric_key"] for metric in metrics}) != EXPECTED_COUNT:
        raise ValueError("V1 指标代码不唯一，停止同步")
    return metrics


def write_csv(metrics: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(metrics[0]))
        writer.writeheader()
        writer.writerows(metrics)


def main() -> None:
    parser = argparse.ArgumentParser(description="同步新疆 V1 黄色指标")
    parser.add_argument("--excel", type=Path, help="指定Excel；不指定时读取搜索目录中最新文件")
    parser.add_argument("--search-dir", type=Path, default=DEFAULT_SEARCH_DIR)
    parser.add_argument("--review", type=Path, default=DEFAULT_REVIEW)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--check-only", action="store_true", help="只识别和验证，不写CSV")
    args = parser.parse_args()
    excel_path = args.excel or find_latest_excel(args.search_dir)
    metrics = extract_metrics(excel_path, args.review)
    for metric in metrics:
        print(
            f"{metric['source_cell']} | {metric['metric_name']} | {metric['value']} "
            f"{metric['unit']} | {metric['year']} | {metric['verification_status']}"
        )
    if not args.check_only:
        write_csv(metrics, args.output)
        print(f"已生成: {args.output}")


if __name__ == "__main__":
    main()
